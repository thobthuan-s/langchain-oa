"""Generate the sample partner customer records workbook.

Usage:
    python samples/partner/create_customer_workbook.py [--owner you@contoso.com] [--contact you@contoso.com]
        [--contact-name "Your Name"]

--owner   receives escalations for every sample account (the account owner).
--contact is added as a verified, billing-authorized Contoso contact so you can
          test from your own mailbox.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

OUTPUT = Path(__file__).with_name("PartnerCustomerRecords.xlsx")


def build(owner: str, contact: str, contact_name: str = "Demo Contact") -> dict[str, list[list[str]]]:
    return {
        "Accounts": [
            ["AccountId", "AccountName", "Domains", "Tier", "AccountOwnerEmail", "SupportSLA", "Status"],
            ["ACC-001", "Contoso Ltd", "contoso.com", "Gold", owner, "Gold: 4-hour response, 24x7", "Active"],
            ["ACC-002", "Fabrikam Inc", "fabrikam.com", "Silver", owner, "Silver: 8 business hours", "Active"],
            ["ACC-003", "Northwind Traders", "northwindtraders.com", "Gold", owner, "Gold: 4-hour response, 24x7", "Active"],
            ["ACC-004", "Tailspin Toys", "tailspintoys.com", "Standard", owner, "Standard: next business day", "Active"],
        ],
        "Contacts": [
            ["ContactEmail", "AccountId", "Name", "Role", "BillingAuthorized"],
            [contact, "ACC-001", contact_name, "IT Director", "Yes"],
            ["megan.bowen@contoso.com", "ACC-001", "Megan Bowen", "IT Administrator", "No"],
            ["alex.wilber@fabrikam.com", "ACC-002", "Alex Wilber", "Finance Manager", "Yes"],
            ["lee.gu@northwindtraders.com", "ACC-003", "Lee Gu", "Operations Lead", "No"],
            ["pat.nguyen@tailspintoys.com", "ACC-004", "Pat Nguyen", "Office Manager", "Yes"],
        ],
        "Subscriptions": [
            ["AccountId", "Product", "Quantity", "AssignedSeats", "Status", "TermEnd", "AutoRenew", "BillingFrequency"],
            ["ACC-001", "Microsoft 365 E5", "250", "238", "Active", "2026-10-20", "No", "Annual"],
            ["ACC-001", "Microsoft 365 Copilot", "50", "50", "Active", "2027-03-31", "Yes", "Annual"],
            ["ACC-001", "Azure plan", "1", "", "Active", "", "", "Monthly"],
            ["ACC-002", "Microsoft 365 Business Premium", "120", "115", "Active", "2027-01-15", "Yes", "Monthly"],
            ["ACC-002", "Dynamics 365 Sales Enterprise", "25", "22", "Active", "2026-12-01", "Yes", "Annual"],
            ["ACC-003", "Microsoft 365 E3", "400", "391", "Active", "2027-06-30", "Yes", "Annual"],
            ["ACC-003", "Microsoft Defender for Endpoint P2", "400", "380", "Active", "2027-06-30", "Yes", "Annual"],
            ["ACC-004", "Microsoft 365 Business Standard", "30", "28", "Active", "2027-04-30", "Yes", "Monthly"],
        ],
        "Cases": [
            ["CaseNumber", "AccountId", "Title", "Status", "Priority", "Opened", "LastUpdated", "NextStep", "InternalNotes"],
            [
                "CON-1017", "ACC-001", "Copilot licenses not appearing for new users", "Resolved", "Normal",
                "2026-09-10", "2026-09-12", "Licenses were re-provisioned; please confirm new users can see Copilot.",
                "Customer escalated twice; exec sponsor is sensitive about delays.",
            ],
            [
                "NW-1042", "ACC-003", "Exchange Online mail flow delays", "In progress", "P1",
                "2026-09-25", "2026-09-27",
                "Microsoft support request 2609250040001234 is open; next update by 2026-09-29 09:00 UTC.",
                "Likely a customer-side connector misconfiguration; do not assign blame in writing.",
            ],
            [
                "NW-1038", "ACC-003", "Teams Rooms device onboarding", "Waiting on customer", "Normal",
                "2026-09-18", "2026-09-24", "Waiting for the device serial numbers from Northwind.", "",
            ],
            [
                "FAB-2203", "ACC-002", "Invoice INV-58213 dispute: duplicate Dynamics charge", "Open", "Normal",
                "2026-09-22", "2026-09-23", "Billing team is reviewing; response by 2026-09-30.",
                "Dispute looks valid; credit of about USD 1,450 pending finance approval.",
            ],
        ],
        "Invoices": [
            ["InvoiceNumber", "AccountId", "InvoiceDate", "Amount", "Currency", "Status", "DueDate"],
            ["INV-58190", "ACC-001", "2026-09-01", "18,750.00", "USD", "Paid", "2026-10-01"],
            ["INV-57544", "ACC-001", "2026-08-01", "18,750.00", "USD", "Paid", "2026-09-01"],
            ["INV-58213", "ACC-002", "2026-09-01", "6,980.00", "USD", "Overdue", "2026-09-16"],
            ["INV-57601", "ACC-002", "2026-08-01", "5,530.00", "USD", "Paid", "2026-08-31"],
            ["INV-58244", "ACC-003", "2026-09-01", "22,400.00", "USD", "Due", "2026-10-01"],
            ["INV-58301", "ACC-004", "2026-09-01", "660.00", "USD", "Paid", "2026-10-01"],
        ],
        "Contracts": [
            ["AccountId", "Agreement", "StartDate", "EndDate", "SLA", "RenewalNoticeDays"],
            ["ACC-001", "Microsoft 365 Managed Services Agreement", "2023-10-21", "2026-10-20", "Gold 24x7, 4-hour response", "60"],
            ["ACC-002", "Cloud Solution Provider Agreement", "2025-01-16", "2027-01-15", "Silver, business hours", "30"],
            ["ACC-003", "Managed Security Services Agreement", "2024-07-01", "2027-06-30", "Gold 24x7, 4-hour response", "90"],
            ["ACC-004", "Cloud Solution Provider Agreement", "2025-05-01", "2027-04-30", "Standard, next business day", "30"],
        ],
        "DemoNotes": [
            ["Scenario", "How to test"],
            ["Contoso renewal (verified, billing)", "Email the agent from the address on the Contacts sheet: 'When does our E5 renew and how many seats are unused?'"],
            ["Northwind P1 case", "Change your Contacts row to AccountId ACC-003, then ask: 'Any update on NW-1042?'"],
            ["Fabrikam billing dispute", "Change your Contacts row to ACC-002, then ask about INV-58213 and request a credit (escalates as a commercial commitment)."],
            ["Unverified sender", "Delete your Contacts row and add your email domain to Tailspin's Domains: the agent routes to the account owner without a draft."],
            ["Hidden data", "Cases.InternalNotes and this sheet are never read by the agent."],
        ],
    }


def write(path: Path, sheets: dict[str, list[list[str]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        for index, column in enumerate(zip(*rows), start=1):
            width = min(60, max(len(str(value)) for value in column) + 2)
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.freeze_panes = "A2"
    workbook.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--owner", default="owner@contoso.com", help="Account owner email for escalations")
    parser.add_argument("--contact", default="demo@contoso.com", help="Your email, added as a Contoso contact")
    parser.add_argument("--contact-name", default="Demo Contact", help="Display name for your contact row")
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    write(args.output, build(args.owner, args.contact, args.contact_name))
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
